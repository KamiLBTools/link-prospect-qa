"""Collect human QA verdicts from the two spreadsheets into one map per domain."""
import openpyxl, re, csv, json, collections

SCALE = {"excellent", "good", "meh", "bad", "manipulation"}

def norm(d):
    d = str(d or "").strip().lower()
    d = re.sub(r"^[a-z]+://", "", d).split("/")[0]
    return d[4:] if d.startswith("www.") else d

def collect():
    entries = collections.defaultdict(list)  # domain -> [(rank, source, verdict, reason, by)]
    wb = openpyxl.load_workbook("/Users/kamiturky/Downloads/QA Training.xlsx", read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True)); hdr = rows[0]
        vi = hdr.index("your_verdict")
        for r in rows[1:]:
            if not r or not r[0]: continue
            d = norm(r[0])
            v = str(r[vi] or "").strip().lower(); why = str(r[vi + 1] or "").strip() if len(r) > vi + 1 else ""
            tail = [x for x in r[vi + 2:] if x not in (None, "") and str(x).strip()]
            second = str(tail[0]).strip().lower() if tail else ""
            second_why = str(tail[1]).strip() if len(tail) > 1 else ""
            reviewer = None if ws.title == "Sheet3" else ws.title
            if second in SCALE:
                first = f" (first review by {reviewer}: {v}{' — ' + why if why else ''})" if v in SCALE else ""
                entries[d].append((1, "QA Training", second, (second_why or "") + first, "Pepo"))
            elif v in SCALE:
                entries[d].append((2 if reviewer else 3, "QA Training", v, why, reviewer or ""))
    wb = openpyxl.load_workbook("/Users/kamiturky/Downloads/New Replies DB.xlsx", read_only=True, data_only=True)
    for ws in wb.worksheets:
        if ws.title == "Sheet3": continue
        rows = list(ws.iter_rows(values_only=True))
        body = rows if ws.title == "LineUP" else rows[1:]  # LineUP has no header row
        for r in body:
            if not r or not r[0] or len(r) < 11: continue
            s = str(r[10] or "").strip().lower()
            v = "good" if s in ("pass",) else "bad" if s in ("fail", "failed", "faill") else None
            if v:
                entries[norm(r[0])].append((4, "Replies DB", v, f"Replies DB qa_status: {str(r[10]).strip()} ({ws.title})", ""))
    return entries

def best(entries):
    out = {}
    for d, es in entries.items():
        es = sorted(es, key=lambda e: e[0])
        top = es[0]
        others = sorted({f"{e[4] or e[1]}: {e[2]}" for e in es[1:] if e[2] != top[2]})
        reason = top[3] + (f" | Other reviews disagree: {'; '.join(others)}" if others else "")
        out[d] = {"QA": top[2], "QA Reason": reason.strip(), "QA By": top[4]}
    return out

if __name__ == "__main__":
    e = collect(); b = best(e)
    json.dump(b, open("qa/human-qa.json", "w"), indent=1)
    res = {r["Website"] for r in csv.DictReader(open("qa/selection-results.csv"))}
    full = {norm(r["website"]) for r in csv.DictReader(open("/Users/kamiturky/Downloads/selection-2026-10-02.csv", encoding="utf-8-sig"))}
    print("domains with a human verdict:", len(b))
    print("in our English results:", len(set(b) & res), "| in the 16,850 list:", len(set(b) & full))
    print("by top source:", collections.Counter(sorted(x, key=lambda y: y[0])[0][1] for x in e.values()))
    print("domains with conflicting reviews:", sum(1 for v in b.values() if "disagree" in v["QA Reason"]))
